"""
Tests for formula-injection protection and folder-name cleaning.

Run with:  python -m pytest tests

These use only small in-memory inputs, so the same tests pass on macOS and
Windows.  They need pytest, openpyxl and pandas but not the GUI libraries.
"""

import csv
import io
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

# Allow running from a fresh checkout without installing the package.
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from markshark.tools.safe_output import (  # noqa: E402
    force_formulas_to_text,
    neutralize_csv_cell,
    neutralize_csv_dataframe,
    neutralize_csv_dict,
    neutralize_csv_row,
)
from markshark.tools.project_utils import sanitize_project_name  # noqa: E402


# ---------------------------------------------------------------- CSV cells
@pytest.mark.parametrize("text", [
    "=1+1", '=HYPERLINK("http://x","y")', "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)",
    "\t=1+1", "\r=1+1",
])
def test_dangerous_text_gets_a_leading_quote(text):
    assert neutralize_csv_cell(text) == "'" + text


@pytest.mark.parametrize("value", [
    "Smith", "O'Brien", "", "A", "-", "=", "A,B", "'=already safe",
    "-5", "+3.2", "-12.5", "1e5",       # plain numbers keep working
    0, 42, 3.5, None,                    # non-text is never changed
])
def test_harmless_values_are_unchanged(value):
    assert neutralize_csv_cell(value) == value


def test_applying_twice_changes_nothing_more():
    once = neutralize_csv_cell("=1+1")
    assert neutralize_csv_cell(once) == once


def test_row_and_dict_helpers():
    assert neutralize_csv_row(["1", "=x1", 5]) == ["1", "'=x1", 5]
    assert neutralize_csv_dict({"LastName": "=x1", "Score": 7}) == {
        "LastName": "'=x1", "Score": 7,
    }


def test_dataframe_helper_keeps_numbers_and_original():
    df = pd.DataFrame({"LastName": ["=evil()", "Smith"], "Score": [-1, 2]})
    safe = neutralize_csv_dataframe(df)
    assert list(safe["LastName"]) == ["'=evil()", "Smith"]
    assert list(safe["Score"]) == [-1, 2]
    assert df["LastName"][0] == "=evil()"  # original untouched


def test_csv_round_trip_shows_quote_not_formula():
    buf = io.StringIO()
    csv.writer(buf).writerow(neutralize_csv_row(["=1+1", "Smith"]))
    assert buf.getvalue().startswith("'=1+1,")


# --------------------------------------------------------------------- XLSX
def test_xlsx_text_beginning_with_equals_is_saved_as_text(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = '=HYPERLINK("http://x","y")'
    ws["A2"] = "Smith"
    ws["A3"] = 5
    assert ws["A1"].data_type == "f"  # what openpyxl does by default

    assert force_formulas_to_text(wb) == 1

    path = tmp_path / "out.xlsx"
    wb.save(path)
    reread = openpyxl.load_workbook(path)  # data_only=False shows formulas
    cell = reread.active["A1"]
    assert cell.data_type == "s"
    assert cell.value == '=HYPERLINK("http://x","y")'
    assert reread.active["A3"].value == 5


# -------------------------------------------------------- folder-name cleaning
@pytest.mark.parametrize("raw, expected", [
    ("Midterm 1", "Midterm 1"),
    ("Test: Spring/Fall", "Test_ Spring_Fall"),
    ("../../etc", "______etc"),
    ("..", "__"),
    ("CON", "CON_"),
    ("aux", "aux_"),
    ("Com3", "Com3_"),
    ("CONSOLE", "CONSOLE"),
    ("   ", ""),
    ("", ""),
])
def test_sanitize_project_name(raw, expected):
    assert sanitize_project_name(raw) == expected


def test_sanitized_name_is_always_a_single_folder():
    for raw in ["../x", "a/b", "a\\b", "C:\\Windows", "/abs", "x/../../y"]:
        name = sanitize_project_name(raw)
        assert "/" not in name and "\\" not in name and ".." not in name


def test_sanitize_limits_length():
    assert len(sanitize_project_name("x" * 500)) == 100
